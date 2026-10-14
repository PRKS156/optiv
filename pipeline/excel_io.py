"""Read the Meridian inventory workbook and write results back without touching provided columns."""
import copy
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import config as C

STUDENT_HEADERS = {  # canonical key -> header prefix in the workbook
    "Criticality Tier": "Criticality Tier",
    "Criticality Rationale": "Criticality Rationale",
    "Assessment Depth Applied": "Assessment Depth Applied",
    "AI Usage Detected": "AI Usage Detected",
    "Evidence Description": "Evidence Description",
    "How Vendor May Be Using AI": "How Vendor May Be Using AI",
    "Named AI Sub-processors": "Named AI Sub-processors",
    "AI Security Risk Classification": "AI Security Risk Classification",
    "Risk Rationale": "Risk Rationale",
    "Recommended Action": "Recommended Action",
    "Assessed": "Assesse",  # template header is spelled 'Assessedd By / Date'
}


def _locate(ws):
    for r in range(1, 15):
        for c in range(1, 30):
            if ws.cell(r, c).value == "Vendor ID":
                return r, c
    raise ValueError("Could not find the 'Vendor ID' header row; is this the Meridian inventory template?")


def read_inventory(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Vendor Inventory"] if "Vendor Inventory" in wb.sheetnames else wb.worksheets[0]
    hr, c0 = _locate(ws)
    headers = {}
    for c in range(c0, ws.max_column + 1):
        v = ws.cell(hr, c).value
        if v:
            headers[c] = str(v).strip()
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        vid = ws.cell(r, c0).value
        if not vid:
            continue
        rec = {h: ws.cell(r, c).value for c, h in headers.items()}
        rec["_row"] = r
        rows.append(rec)
    return rows


def write_results(src_path, out_path, results, team="Team", skip_ids=("V-000",), config_snapshot=True):
    wb = openpyxl.load_workbook(src_path)
    ws = wb["Vendor Inventory"] if "Vendor Inventory" in wb.sheetnames else wb.worksheets[0]
    hr, c0 = _locate(ws)
    col = {}
    for c in range(c0, ws.max_column + 1):
        h = str(ws.cell(hr, c).value or "")
        for key, prefix in STUDENT_HEADERS.items():
            if h.startswith(prefix):
                col[key] = c
    inv = {str(ws.cell(r, c0).value): r for r in range(hr + 1, ws.max_row + 1) if ws.cell(r, c0).value}
    model_row = inv.get("V-000")

    for res in results:
        vid = res["vendor_id"]
        if vid in skip_ids or vid not in inv:
            continue
        r = inv[vid]
        for key, val in res["cells"].items():
            if key not in col:
                continue
            cell = ws.cell(r, col[key])
            cell.value = val.replace("{team}", team) if isinstance(val, str) else val
            if model_row:  # inherit the template's student-cell formatting
                src = ws.cell(model_row, col[key])
                cell._style = copy.copy(src._style)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(name, headers, rows, widths=None):
        if name in wb.sheetnames:
            del wb[name]
        s = wb.create_sheet(name)
        s.append(headers)
        for c in s[1]:
            c.fill, c.font = head_fill, head_font
        for row in rows:
            s.append(row)
        for i, h in enumerate(headers):
            s.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = (widths or {}).get(i, 22)
        for r_ in s.iter_rows(min_row=2):
            for c in r_:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        s.freeze_panes = "A2"
        return s

    # Evidence Log: one row per signal, fully traceable
    ev_rows = []
    for res in results:
        for s in res["signals"]:
            ev_rows.append([res["vendor_id"], res["vendor_name"], s["type"], s["excerpt"], s["source_class"],
                            s["source_weight"], s["trust"], s["score"], s["url"], s["retrieved"], s["sha256"],
                            s["provenance"], "; ".join(s["providers"]), "yes" if s["llm_assisted"] else "no"])
        for e in res["errors"]:
            ev_rows.append([res["vendor_id"], res["vendor_name"], "FETCH_NOTE", e, "", "", "", "", "", "", "", "", "", ""])
        for d in res["docs"]:
            ev_rows.append([res["vendor_id"], res["vendor_name"], "DOCUMENT_REVIEWED", d.get("title", ""),
                            d["source_class"], "", "", "", d["url"], d["retrieved"], d["sha256"][:16],
                            d.get("provenance", ""), "", ""])
    sheet("Evidence Log",
          ["Vendor ID", "Vendor", "Signal type", "Exact excerpt / note", "Source class", "Source weight",
           "Trust factor", "Signal score", "URL", "Retrieved", "Content SHA-256 (16)", "Provenance",
           "Providers named", "LLM assisted"], ev_rows, {3: 70, 8: 45})

    sheet("Criticality Scores",
          ["Vendor ID", "Vendor", "Data sensitivity", "Operational dependency", "Payment flow",
           "Regulatory exposure", "Data volume", "Weighted score", "Tier", "Floors applied",
           "AI usage score", "Evidence tag", "AI usage", "Severity", "AI risk class", "Flags"],
          [[r["vendor_id"], r["vendor_name"], *r["crit"]["factors"].values(), r["crit"]["weighted_score"],
            r["crit"]["tier"], " | ".join(r["crit"]["floors_applied"]), r["weighed"]["usage_score"],
            r["weighed"]["tag"], r["weighed"]["detected"], r["risk"]["severity"], r["risk"]["risk_class"],
            ", ".join(r["risk"]["flags"])] for r in results], {1: 30, 9: 50, 15: 40})

    if config_snapshot:
        rub = [["Criticality factor weight", k, v, ""] for k, v in C.CRITICALITY_WEIGHTS.items()]
        rub += [["Tier threshold (score >=)", t, th, ""] for t, th in C.TIER_THRESHOLDS]
        rub += [["Source class weight", k, w, why] for k, (w, why) in C.SOURCE_CLASSES.items()]
        rub += [["Signal type multiplier", k, v, ""] for k, v in C.SIGNAL_TYPES.items()]
        rub += [["Evidence tag threshold (usage score >=)", t, th, ""] for t, th in C.TAG_RULES]
        rub += [["Risk threshold (severity >=)", t, th, ""] for t, th in C.RISK_THRESHOLDS]
        rub += [["Depth: " + t, d["label"], d["max_pages"], d["stop_rule"]] for t, d in C.DEPTH_MATRIX.items()]
        sheet("Rubric & Weights", ["Parameter", "Name", "Value", "Why / rule"], rub, {0: 38, 1: 26, 3: 90})

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
